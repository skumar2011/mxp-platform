#!/usr/bin/env python3
"""Build the MXP requirements, table-design and rules workbook from the repo sources."""

import json
import math
import re
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = ROOT / "packages" / "engine" / "src" / "schema.sql"
RULES_DIR = ROOT / "samples" / "rules"
CAMPAIGNS = ROOT / "samples" / "campaigns.json"
OUT = ROOT / "docs" / "MXP-requirements-schema-rules.xlsx"

NAVY = "1B3A4B"
ALT = "F4F7FA"
MUST = "E5F6EA"
SHOULD = "FFF6D8"
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, color=NAVY, size=20)
SECTION_FONT = Font(name="Calibri", bold=True, color=NAVY, size=13)
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
LABEL_FONT = Font(name="Calibri", bold=True, size=11, color="1B3A4B")

REQ = "1B4F72"
DATA = "0E6655"
RULE = "9A7D0A"


def load_db():
    conn = sqlite3.connect(":memory:")
    conn.executescript(SCHEMA.read_text())
    return conn


def extract_checks(sql: str):
    checks = []
    i = 0
    upper = sql.upper()
    while True:
        j = upper.find("CHECK", i)
        if j < 0:
            break
        k = sql.find("(", j)
        if k < 0:
            break
        depth = 0
        end = None
        for p in range(k, len(sql)):
            if sql[p] == "(":
                depth += 1
            elif sql[p] == ")":
                depth -= 1
                if depth == 0:
                    end = p
                    break
        if end is None:
            break
        checks.append(re.sub(r"\s+", " ", sql[k + 1 : end].strip()))
        i = end + 1
    return checks


def column_comments(sql_text: str):
    comments = {}
    for match in re.finditer(r"CREATE TABLE\s+(\w+)\s*\((.*?)\);", sql_text, re.S):
        table = match.group(1)
        body = match.group(2)
        for line in body.splitlines():
            raw = line.strip().rstrip(",")
            if not raw or raw.startswith("--") or raw.upper().startswith(
                ("PRIMARY", "UNIQUE", "CHECK", "FOREIGN", "CONSTRAINT")
            ):
                continue
            col = raw.split()[0]
            if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", col):
                continue
            cmt = ""
            if "--" in line:
                cmt = line.split("--", 1)[1].strip()
            if cmt:
                comments[(table, col)] = cmt
    return comments


TABLES = {
    "channel_ref": ("Reference", "Reference", "Channel codes validated on ingest and offered in the rule builder."),
    "product_ref": ("Reference", "Reference", "Product codes and families (SAVINGS, CREDIT_CARD, UPI, and so on)."),
    "txn_type_ref": ("Reference", "Reference", "Transaction types. is_reversal marks reversal types."),
    "txn_category_ref": ("Reference", "Reference", "Transaction categories such as RENT, SELF_TRANSFER, UTILITY_BILL."),
    "segment": ("Reference", "Reference", "Customer segments used by customer.segments conditions."),
    "customer": ("Customer", "Master", "Tokenised customer. Masked name only; no PAN, Aadhaar or raw account numbers."),
    "account": ("Customer", "Master", "Savings, card, UPI and other accounts. account_ref is a token."),
    "instrument": ("Customer", "Master", "Card, VPA or wallet identifier. instrument_ref is a token, never a PAN."),
    "customer_segment": ("Customer", "Master", "Segment membership with validity dates."),
    "ingestion_batch": ("Ingestion", "Operational", "File or batch with control totals. Status OPEN → CLOSED → RECONCILED or MISMATCH."),
    "ingestion_event": ("Ingestion", "Operational", "Every delivery, stored verbatim so it can be replayed."),
    "ingestion_event_history": ("Ingestion", "Fact", "Per-stage trail: INGEST, NORMALIZE, EVALUATE, RETRY, REPLAY, DEAD_LETTER."),
    "txn_master": ("Transactions", "Master", "Central transaction master. Financial fields are frozen by trigger and rows are never deleted."),
    "txn_status_history": ("Transactions", "Fact", "Forward-only status log, including refunded amount and the causing event."),
    "balance_snapshot": ("Transactions", "Master", "Daily closing balances. Never inferred from transactions. Revisions are kept."),
    "campaign": ("Rules", "Configuration", "Campaign window, budget and enrolment mode."),
    "rule": ("Rules", "Configuration", "Stable rule identity. rule_code is unique."),
    "rule_version": ("Rules", "Configuration", "One version. Immutable once submitted. At most one ACTIVE or PAUSED version per rule."),
    "rule_condition_group": ("Rules", "Configuration", "Eligibility or qualifying group. AND/OR, optional negation, nesting."),
    "rule_condition": ("Rules", "Configuration", "One comparison: field, one of 13 operators, value_json."),
    "rule_tier": ("Rules", "Configuration", "Points ladder. Threshold is in rupees for money metrics, or a count."),
    "rule_component": ("Rules", "Configuration", "Component goals of a COMBINED rule, with a minimum tier."),
    "customer_enrollment": ("Rules", "Operational", "Campaign enrolment: AUTO, OPT_IN or ADMIN; ENROLLED or OPTED_OUT."),
    "goal_progress": ("Progress", "Projection", "Progress for one customer, rule version and period. Rebuildable from the journal."),
    "goal_contribution": ("Progress", "Fact", "Append-only journal of APPLY and REVERSE deltas behind goal_progress."),
    "rule_evaluation": ("Progress", "Operational", "One row per evaluation: outcome, reason and full trace."),
    "xp_ledger": ("Rewards", "Fact", "Immutable XP ledger. AWARD, REVERSAL or ADJUSTMENT. Idempotency key is unique."),
    "xp_account": ("Rewards", "Projection", "Balance, lifetime earned and lifetime reversed. Reconciled against the ledger."),
    "reward_adjustment": ("Rewards", "Operational", "Maker-checker manual adjustment. REQUESTED → APPROVED or REJECTED."),
    "nudge_template": ("Nudges", "Configuration", "Message template per nudge type and channel."),
    "nudge": ("Nudges", "Operational", "Generated nudge. dedupe_key is unique. Suppressed rows record why."),
    "app_user": ("Administration", "Administration", "Console identity. Password stored as a scrypt hash. Lock-out counters."),
    "user_role": ("Administration", "Administration", "Roles assigned to a console user."),
    "api_client": ("Administration", "Administration", "Source-system API key. Only key_hash is stored."),
    "system_config": ("Administration", "Administration", "Engine settings. Key engine.config."),
    "audit_log": ("Administration", "Fact", "Hash-chained, append-only audit of who changed what, when and why."),
    "reconciliation_run": ("Administration", "Operational", "One reconciliation run and its pass or fail result."),
    "reconciliation_break": ("Administration", "Operational", "A break found by a run: expected versus actual."),
}

INDEX_PURPOSE = {
    "ix_account_customer": "Accounts of a customer",
    "ix_instrument_customer": "Instruments of a customer",
    "ix_ing_queue": "Worker queue: status and next attempt time",
    "ix_ing_txn": "Event resolved to a transaction",
    "ix_ing_batch": "Events in a batch",
    "ix_ing_hist_event": "History of one ingestion event",
    "ix_txn_customer_ts": "Customer history by time",
    "ix_txn_account": "Account history by time",
    "ix_txn_original": "Refund or reversal back to the original",
    "ix_txn_status": "Transactions by status",
    "ix_txn_proc": "Transactions awaiting evaluation",
    "ix_txn_hist": "Status history of one transaction",
    "ix_bal_account_date": "Balance window for one account",
    "ix_rv_status": "Rule versions by lifecycle status",
    "ix_rv_campaign": "Rule versions in a campaign",
    "ux_rv_one_live": "At most one ACTIVE or PAUSED version per rule",
    "ix_rcg_version": "Condition groups of a version",
    "ix_rc_group": "Conditions in a group",
    "ix_gp_rule": "Progress for a rule and period",
    "ix_gp_open": "Open progress rows due for period close",
    "ix_gc_txn": "Journal rows for a transaction",
    "ix_gc_progress": "Journal behind one progress row",
    "ix_re_customer": "Evaluations for a customer",
    "ix_re_rule": "Evaluations for a rule",
    "ix_re_txn": "Evaluations for a transaction",
    "ix_re_outcome": "Evaluations by outcome",
    "ix_xl_customer": "Customer ledger in order",
    "ix_xl_slot": "Net points of one reward slot",
    "ix_xl_rule": "Ledger rows of a rule",
    "ix_xl_campaign": "Ledger rows of a campaign",
    "ix_xl_reverses": "Reversal pointing at an award",
    "ix_xl_txn": "Ledger rows of a source transaction",
    "ix_nudge_customer": "Nudges for a customer",
    "ix_nudge_status": "Nudges by delivery status",
    "ix_audit_entity": "Audit rows for an entity",
    "ix_audit_time": "Audit rows by time",
    "ix_rb_run": "Breaks of a reconciliation run",
}

REL_NOTE = {
    ("account", "customer_id"): "A customer owns many accounts",
    ("instrument", "customer_id"): "A customer holds many instruments",
    ("instrument", "account_id"): "An instrument may be linked to one account",
    ("customer_segment", "customer_id"): "A customer belongs to many segments",
    ("customer_segment", "segment_code"): "A segment has many members",
    ("account", "product_code"): "An account is typed by one product",
    ("txn_master", "customer_id"): "A customer makes many transactions",
    ("txn_master", "account_id"): "An account has many transactions",
    ("txn_master", "original_txn_id"): "A refund or reversal points at the original transaction",
    ("txn_status_history", "txn_id"): "A transaction has many status changes",
    ("balance_snapshot", "account_id"): "An account has one closing snapshot per day",
    ("ingestion_event", "batch_id"): "A batch contains many events",
    ("ingestion_event_history", "event_id"): "An event has many stage-log rows",
    ("rule_version", "rule_id"): "A rule has many versions",
    ("rule_version", "campaign_id"): "A campaign groups many rule versions",
    ("rule_condition_group", "rule_version_id"): "A version has eligibility and qualifying groups",
    ("rule_condition_group", "parent_group_id"): "A group may nest inside another group",
    ("rule_condition", "group_id"): "A group contains many conditions",
    ("rule_tier", "rule_version_id"): "A version has a points ladder",
    ("rule_component", "rule_version_id"): "A combined version lists component rules",
    ("rule_component", "component_rule_id"): "A rule may be a component of combined rules",
    ("goal_progress", "customer_id"): "Progress is kept per customer, version and period",
    ("goal_progress", "rule_version_id"): "Progress belongs to the version that produced it",
    ("goal_contribution", "txn_id"): "A transaction contributes to many rules",
    ("rule_evaluation", "customer_id"): "Evaluations are stored per customer",
    ("xp_ledger", "customer_id"): "A customer earns many ledger rows",
    ("xp_ledger", "reverses_ledger_id"): "A reversal must point at the award it undoes",
    ("xp_account", "customer_id"): "One balance projection per customer",
    ("nudge", "customer_id"): "A customer receives many nudges",
    ("audit_log", "actor"): "Actor is stored by name, not as a foreign key",
}


def write_table(ws, headers, rows, widths, wrap=True, priority_col=None):
    header_fill = PatternFill("solid", fgColor=NAVY)
    alt = PatternFill("solid", fgColor=ALT)
    must = PatternFill("solid", fgColor=MUST)
    should = PatternFill("solid", fgColor=SHOULD)
    for c, header in enumerate(headers, 1):
        cell = ws.cell(1, c, header)
        cell.fill = header_fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        cell.border = THIN
    ws.row_dimensions[1].height = 24
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for r, row in enumerate(rows, 2):
        lines = 1
        for c, value in enumerate(row, 1):
            text = "" if value is None else value
            cell = ws.cell(r, c, text)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")
            cell.border = THIN
            if r % 2 == 0:
                cell.fill = alt
            if priority_col and c == priority_col:
                if text == "M":
                    cell.fill = must
                elif text == "S":
                    cell.fill = should
            width = widths[c - 1]
            chars = max(8, int(width * 1.15))
            n = 0
            for part in str(text).split("\n"):
                n += max(1, math.ceil(len(part) / chars)) if part else 1
            lines = max(lines, n)
        ws.row_dimensions[r].height = min(96, max(18, lines * 15))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.7, bottom=0.6, header=0.3, footer=0.3)
    ws.oddHeader.left.text = "Money XP (MXP)"
    ws.oddHeader.right.text = "Phase 1"
    ws.oddFooter.left.text = ws.title
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = "1:1"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.pageOrder = "downThenOver"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def add_sheet(wb, name, color, headers, rows, widths, priority_col=None):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = color
    write_table(ws, headers, rows, widths, priority_col=priority_col)
    return ws


def flatten_conditions(defn):
    rows = []

    def walk(node, purpose, path, group_op, negated):
        if not node:
            return
        if node.get("type") == "group":
            op = node.get("op", "AND")
            neg = bool(node.get("negate"))
            label = f"NOT ({op})" if neg else op
            next_path = f"{path} / {label}" if path else label
            for child in node.get("children", []):
                walk(child, purpose, next_path, op, neg)
            return
        value = node.get("value")
        if isinstance(value, list):
            shown = ", ".join(str(v) for v in value)
        elif value is None:
            shown = ""
        else:
            shown = str(value)
        op = node["operator"]
        field = node["field"]
        phrase = {
            "EQ": f"{field} equals {shown}",
            "NE": f"{field} does not equal {shown}",
            "GT": f"{field} is greater than {shown}",
            "GTE": f"{field} is greater than or equal to {shown}",
            "LT": f"{field} is less than {shown}",
            "LTE": f"{field} is less than or equal to {shown}",
            "IN": f"{field} is one of {shown}",
            "NOT_IN": f"{field} is not one of {shown}",
            "BETWEEN": f"{field} is between {shown}",
            "CONTAINS": f"{field} contains {shown}",
            "STARTS_WITH": f"{field} starts with {shown}",
            "IS_NULL": f"{field} is empty",
            "NOT_NULL": f"{field} is present",
        }.get(op, f"{field} {op} {shown}")
        if negated:
            phrase = "NOT ( " + phrase + " )"
        rows.append((purpose, path, "Yes" if negated else "No", field, op, shown, phrase))

    for purpose in ("eligibility", "qualifying"):
        walk(defn.get(purpose), purpose.upper(), "", "", False)
    return rows


def main():
    conn = load_db()
    comments = column_comments(SCHEMA.read_text())
    master = {
        name: sql
        for name, sql in conn.execute(
            "SELECT name, sql FROM sqlite_master WHERE sql IS NOT NULL"
        )
    }
    table_names = [
        name
        for (name,) in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY rowid"
        )
    ]

    wb = Workbook()
    cover = wb.active
    cover.title = "Contents"

    # ----- Actors -----
    add_sheet(
        wb,
        "Actors",
        REQ,
        ["Actor", "Role code", "Needs"],
        [
            ["Product / campaign owner", "RULE_MAKER", "Create and edit campaigns and rule drafts; run simulations"],
            ["Risk / business approver", "RULE_APPROVER", "Review, approve or reject; approve manual adjustments"],
            ["QA engineer", "QA_ENGINEER", "Build scenarios, matrix runs and regression suites in the Testing app"],
            ["Operations", "OPERATIONS", "Ingestion health, replay, reconciliation, activate / pause, request adjustments, exports"],
            ["Auditor", "AUDITOR", "Read-only access to everything including the audit log; exports"],
            ["System administrator", "SYSTEM_ADMIN", "Users, API keys, engine configuration, customer import"],
            ["Source systems", "API key", "Deliver transactions and balance snapshots"],
            ["Customer", "—", "Receives nudges and sees XP. The customer-facing app is out of scope for Phase 1"],
        ],
        [32, 22, 88],
    )

    # ----- Use cases -----
    add_sheet(
        wb,
        "Use Cases",
        REQ,
        ["#", "Use case", "How it is expressed", "Sample rule code"],
        [
            ["1", "Average monthly savings balance ≥ ₹50,000 → 500 XP", "GOAL, metric AVG_BALANCE, monthly, one tier", "SAV_AVG_BAL_50K"],
            ["2", "Card spend ≥ ₹10,000 in a month → 300 XP, plus a higher tier at ₹25,000 → 400 XP", "GOAL, metric SUM, qualifying = card purchases, cumulative tiers", "CARD_SPEND_10K"],
            ["3", "10 eligible UPI payments in a month → 100 XP", "GOAL, metric COUNT. Self-transfers and amounts under ₹10 are excluded", "UPI_10_TXNS"],
            ["4", "Three utility bill payments in a campaign period → bonus; six bills → a higher tier", "GOAL, COUNT, period CAMPAIGN, tiers 3 → 150 and 6 → 250", "UTILITY_3_BILLS"],
            ["5", "Savings growth of ₹25,000 → 400 XP", "GOAL, metric BALANCE_GROWTH, monthly", "SAV_GROWTH_25K"],
            ["6", "Combined goals (balance and card spend) → 1,000 XP", "COMBINED of use cases 1 and 2, both at tier ≥ 1", "BALANCE_PLUS_CARD"],
            ["7", "1 XP per ₹100 on merchant UPI, capped at 200 XP a month", "PER_EVENT, PER_AMOUNT, monthly cap. Left pending approval in the sample load", "UPI_MERCHANT_PER_TXN"],
        ],
        [8, 78, 62, 28],
    )

    functional = [
        ["F-01", "Transaction master and ingestion", "M", "One canonical transaction master for all channels; extensible attributes without a schema change (attributes_json)."],
        ["F-02", "Transaction master and ingestion", "M", "Source adapters deliver events in a canonical envelope (TRANSACTION, BALANCE_SNAPSHOT) via API or batch."],
        ["F-03", "Transaction master and ingestion", "M", "Validate (schema, reference data, customer/account/instrument existence, future timestamps), normalise (money to minor units, timestamps to UTC with a business date in the rule time zone) and de-duplicate on (source_system, source_event_id) and (source_system, source_txn_id)."],
        ["F-04", "Transaction master and ingestion", "M", "The same event id with different content is a CONFLICTING_DUPLICATE and is rejected, not applied."],
        ["F-05", "Transaction master and ingestion", "M", "Unknown customer or account is retried with exponential back-off (SUSPENDED), then dead-lettered after maxAttempts. Every stage is recorded in the event history and can be replayed."],
        ["F-06", "Transaction master and ingestion", "M", "Status transitions are forward-only: pending → posted → reversed / refunded / partially refunded; failed. Refunds and reversals link to the original transaction."],
        ["F-07", "Transaction master and ingestion", "M", "Batch files carry control totals. Closing a batch reports RECONCILED or MISMATCH."],
        ["F-08", "Transaction master and ingestion", "M", "Balances come from daily closing balance snapshots, never inferred from transactions. Revisions are kept."],
        ["F-10", "Rules engine", "M", "Rule kinds: GOAL (count, sum, average balance, minimum balance, balance growth), PER_EVENT (fixed or per-amount points), COMBINED (number of component goals met)."],
        ["F-11", "Rules engine", "M", "Condition trees with nested AND / OR groups and negation over txn.*, account.*, customer.* and txn.attr.<key>. Thirteen operators. Separate eligibility (who) and qualifying (what counts) groups."],
        ["F-12", "Rules engine", "M", "Periods: daily, weekly, monthly, quarterly, campaign. Time zone per rule. Transaction date or posting date basis. Late-arrival grace days."],
        ["F-13", "Rules engine", "M", "Tiers with CUMULATIVE (each tier adds) or HIGHEST_ONLY modes."],
        ["F-14", "Rules engine", "M", "Limits: once or recurring, max occurrences, points cap per period, lifetime cap, campaign budget."],
        ["F-15", "Rules engine", "M", "Priority and conflict groups with policies CUMULATIVE, EXCLUSIVE_HIGHEST_POINTS, EXCLUSIVE_PRIORITY."],
        ["F-16", "Rules engine", "M", "Reversal handling per rule: CLAWBACK (withdraw points when a qualifying transaction is reversed or refunded below the threshold) or GRANDFATHER."],
        ["F-17", "Rules engine", "M", "Which transaction statuses count (POSTED, optionally PENDING). FAILED never counts."],
        ["F-18", "Rules engine", "M", "Effective dates on rules and campaigns. Transactions outside the window never count."],
        ["F-19", "Rules engine", "M", "Idempotent rewards: the same facts processed any number of times, in any order, yield exactly the same ledger."],
        ["F-20", "Rules engine", "M", "Honest data gaps: a balance goal with missing snapshots reports DATA_UNAVAILABLE rather than guessing."],
        ["F-30", "Rewards", "M", "Append-only XP ledger with AWARD, REVERSAL, ADJUSTMENT. Each entry carries rule version, period, reason, idempotency key and links to the triggering evaluation and transaction."],
        ["F-31", "Rewards", "M", "Customer XP account (balance, lifetime earned, lifetime reversed) is a projection reconciled against the ledger."],
        ["F-32", "Rewards", "M", "Manual adjustments require a reason and a second person (maker-checker)."],
        ["F-33", "Rewards", "M", "Rewards for a closed period are never created by late data. The metric is still updated."],
        ["F-40", "Nudges", "M", "Generate PROGRESS, NEAR_TARGET, TIER_ACHIEVED and REWARD_REVERSED nudges from goal progress."],
        ["F-41", "Nudges", "M", "Only approved channels, only with customer consent, daily cap per customer, de-duplicated. Suppressed nudges record why."],
        ["F-42", "Nudges", "S", "Delivery goes through an adapter. Phase 1 ships a logging adapter."],
        ["F-50", "Admin Console", "M", "Dashboard: awards, reversals, active rules, pending approvals, ingestion health, reconciliation status."],
        ["F-51", "Admin Console", "M", "Rule builder with live validation messages, condition-tree editor, tier editor, and a plain-language description."],
        ["F-52", "Admin Console", "M", "Lifecycle: draft → validate → simulate → submit → approve → activate → pause / resume → retire; reject; roll back."],
        ["F-53", "Admin Console", "M", "Version history, side-by-side diff, rollback with reason."],
        ["F-54", "Admin Console", "M", "Maker-checker: the person who created or submitted a version can never approve it."],
        ["F-55", "Admin Console", "M", "RBAC with the six roles. Every state change is audited."],
        ["F-56", "Admin Console", "M", "Evaluation search with full trace, ledger search, customer 360, CSV export (audited)."],
        ["F-57", "Admin Console", "M", "Campaigns, ingestion monitor (events, history, replay, batches), reconciliation, audit viewer with chain verification, user / API-key / configuration administration."],
        ["F-60", "Testing and Simulation", "M", "Separate responsive app. Select rule, campaign, version (including an unsaved draft) and a synthetic or real-history customer."],
        ["F-61", "Testing and Simulation", "M", "Scenario language covering transactions, status changes, reversals, refunds, balance snapshots and period close; duplicate deliveries; overlapping rules."],
        ["F-62", "Testing and Simulation", "M", "Single-rule and multi-rule runs; auto-generated scenario matrix (positive, negative, boundary, eligibility, duplicate, reversal, refund, status, period, regression); performance runs."],
        ["F-63", "Testing and Simulation", "M", "Save, clone, rerun cases; suites with history; expected versus actual per check."],
        ["F-64", "Testing and Simulation", "M", "Per run: inputs, rules and versions, per-condition pass/fail with reasons, XP, progress before/after, caps, exclusions, conflicts, full trace, errors, downloadable CSV."],
        ["F-65", "Testing and Simulation", "M", "Production isolation: simulations run in a throw-away in-memory database with the same engine code. Production data is read-only. Results are stored in a separate simulation database."],
        ["F-70", "Operations and compliance", "M", "Reconciliation: ledger vs XP account, progress vs journal, journal vs master, reversal integrity, stuck pipeline, dead letters, batch control, campaign budget, audit chain; optional source reconciliation."],
        ["F-71", "Operations and compliance", "M", "Period close, rule backfill and event replay are idempotent operator actions."],
        ["F-72", "Operations and compliance", "M", "Immutable audit log, hash-chained and tamper-evident."],
    ]
    add_sheet(
        wb,
        "Functional Requirements",
        REQ,
        ["ID", "Area", "Priority", "Requirement"],
        functional,
        [10, 36, 12, 110],
        priority_col=3,
    )

    add_sheet(
        wb,
        "Non-Functional",
        REQ,
        ["Area", "Requirement", "How it is met"],
        [
            ["Correctness", "No duplicate or lost rewards", "Idempotency keys, convergent award algorithm, forward-only transitions, automated tests"],
            ["Performance", "Never scan the whole transaction table per rule", "Transaction-driven evaluation, incremental goal_progress, append-only journal"],
            ["Security", "Authentication, authorisation, encryption, masking, secrets", "See deployment and security runbooks (docs/10)"],
            ["Auditability", "Who changed what, when, and why", "Hash-chained audit log; sealed rule hashes"],
            ["Availability", "Retry, dead-letter queue, restartable worker, disaster recovery", "Event store is the queue; replay; runbooks"],
            ["Observability", "Health, readiness, reconciliation, structured logs", "/api/health, /api/ready, /api/ops/health, correlation ids"],
            ["Environments", "Dev / test / staging / prod separation", "MXP_ENV, no demo data in prod, strong JWT secret enforced"],
            ["Change control", "Controlled activation and rollback", "Maker-checker, one open version per rule, rollback"],
            ["Privacy", "Data minimisation", "Tokens and masked names only; no PAN, account numbers or raw PII"],
        ],
        [22, 62, 78],
    )

    add_sheet(
        wb,
        "Out of Scope",
        REQ,
        ["Item", "Phase 1 position", "Seam left in the design"],
        [
            ["Customer-facing app and redemption / catalogue", "Out of scope", "XP ledger and account are the source a later app would read"],
            ["Real SMS / push providers", "Out of scope", "Nudge delivery is an adapter. Phase 1 ships a logging adapter (F-42)"],
            ["SSO / identity-provider integration", "Out of scope", "Local login with roles; an IdP can replace it"],
            ["Multi-currency", "Out of scope", "Currency is stored (default INR) but evaluation assumes one currency"],
            ["Streaming infrastructure (Kafka)", "Out of scope", "Ingestion is API or batch into ingestion_event, which is the queue"],
        ],
        [52, 22, 78],
    )

    add_sheet(
        wb,
        "Conventions",
        DATA,
        ["Topic", "Rule"],
        [
            ["Money", "Stored as integer minor units (paise). Rule thresholds are written in rupees and converted on load. Never floats."],
            ["Timestamps", "ISO-8601 UTC text. Business dates are YYYY-MM-DD in the rule's time zone."],
            ["Facts", "Append-only: xp_ledger, goal_contribution, audit_log, txn_status_history, ingestion_event_history. Update and delete are blocked by triggers where stated."],
            ["Projections", "goal_progress and xp_account can be rebuilt from facts and are reconciled."],
            ["Transaction master", "Financial source of truth. MXP never writes gamification outcomes into it and never posts to source systems."],
            ["Privacy", "customer_ref and account_ref are tokens. The customer name is masked. No PAN, Aadhaar or raw account numbers are stored."],
            ["Rule storage", "A version is stored as normalised rows and as a sealed JSON definition. The hash is verified on every load. A mismatch raises DEFINITION_TAMPERED."],
            ["Parameters", "Thresholds, periods, tiers, points, products, channels, categories, exclusions, caps, frequency, dates, priority and conflict handling are all data. None of the use cases are hard-coded."],
            ["Acceptance", "An admin creates a goal, configures conditions and XP, tests it against scenarios, has it approved and activated, and verifies correct rewards with no duplicates."],
        ],
        [24, 120],
    )

    catalog_rows = []
    for name in table_names:
        domain, kind, purpose = TABLES.get(name, ("Other", "Operational", ""))
        catalog_rows.append([name, domain, kind, purpose])
    add_sheet(
        wb,
        "Table Catalog",
        DATA,
        ["Table", "Domain", "Kind", "Purpose"],
        catalog_rows,
        [32, 20, 18, 100],
    )

    column_rows = []
    for table in table_names:
        info = conn.execute(f"PRAGMA table_info({table})").fetchall()
        fks = conn.execute(f"PRAGMA foreign_key_list({table})").fetchall()
        fk_map = {}
        for fk in fks:
            # id, seq, table, from, to, on_update, on_delete, match
            fk_map.setdefault(fk[3], []).append(f"{fk[2]}({fk[4]})")
        checks = extract_checks(master.get(table, ""))
        uniq = {}
        index_sql = {
            name: sql
            for name, sql in conn.execute(
                "SELECT name, sql FROM sqlite_master WHERE type='index' AND tbl_name=?",
                (table,),
            )
        }
        for _seq, idx_name, is_unique, origin, partial in conn.execute(f"PRAGMA index_list({table})"):
            if not is_unique:
                continue
            cols = [row[2] for row in conn.execute(f"PRAGMA index_info('{idx_name}')")]
            where = ""
            if partial and index_sql.get(idx_name):
                wm = re.search(r"WHERE\s+(.+)$", index_sql[idx_name], re.I)
                if wm:
                    where = " WHERE " + re.sub(r"\s+", " ", wm.group(1).strip())
            if origin == "pk":
                continue
            if len(cols) == 1 and not where:
                label = "Yes"
            else:
                label = "UNIQUE (" + ", ".join(cols) + ")" + where
            for col in cols:
                uniq.setdefault(col, [])
                if label not in uniq[col]:
                    uniq[col].append(label)
        domain = TABLES.get(table, ("", "", ""))[0]
        for cid, col, ctype, notnull, default, pk in info:
            col_checks = []
            for check in checks:
                if re.search(rf"\b{re.escape(col)}\b", check) and check.count(",") < 12:
                    # skip huge multi-column checks that only mention the column in passing? keep if the check is about this column
                    names = set(re.findall(r"\b[a-z_][a-z0-9_]*\b", check))
                    keywords = {"and", "or", "in", "is", "null", "not", "between", "check", "case", "when", "then", "else", "end"}
                    cols_mentioned = [n for n in names if n not in keywords and not n.isupper() and n != col]
                    # values inside quotes are uppercase typically. Column names are lowercase.
                    if len(cols_mentioned) <= 1:
                        col_checks.append(check)
            refs = ", ".join(fk_map.get(col, []))
            unique = "; ".join(uniq.get(col, []))
            column_rows.append(
                [
                    domain,
                    table,
                    col,
                    ctype,
                    "Yes" if pk else "",
                    "Yes" if notnull or pk else "",
                    "" if default is None else str(default),
                    refs,
                    unique,
                    "; ".join(col_checks),
                    comments.get((table, col), ""),
                ]
            )
    add_sheet(
        wb,
        "Column Design",
        DATA,
        ["Domain", "Table", "Column", "Data type", "Primary key", "Required", "Default", "References", "Unique / key", "Check", "Comment"],
        column_rows,
        [16, 28, 26, 14, 14, 12, 22, 36, 42, 55, 62],
    )

    rel_rows = []
    for table in table_names:
        for fk in conn.execute(f"PRAGMA foreign_key_list({table})"):
            note = REL_NOTE.get((table, fk[3]), "Many-to-one")
            rel_rows.append([table, fk[3], fk[2], fk[4], fk[6] or "NO ACTION", fk[5] or "NO ACTION", note])
    rel_rows.append(["audit_log", "actor (name, not a foreign key)", "app_user", "username", "—", "—", REL_NOTE[("audit_log", "actor")]])
    add_sheet(
        wb,
        "Relationships",
        DATA,
        ["From table", "From column", "To table", "To column", "On delete", "On update", "Relationship"],
        rel_rows,
        [28, 36, 28, 22, 16, 16, 62],
    )

    index_rows = []
    for name, tbl, sql in conn.execute(
        "SELECT name, tbl_name, sql FROM sqlite_master WHERE type='index' AND sql IS NOT NULL ORDER BY tbl_name, name"
    ):
        unique = "Yes" if "UNIQUE" in sql.upper() else ""
        m = re.search(r"\(([^)]*)\)", sql)
        cols = m.group(1).replace(" ", "") if m else ""
        where = ""
        wm = re.search(r"WHERE\s+(.+)$", sql, re.I)
        if wm:
            where = re.sub(r"\s+", " ", wm.group(1).strip())
        index_rows.append([name, tbl, unique, cols, where, INDEX_PURPOSE.get(name, "")])
    add_sheet(
        wb,
        "Indexes",
        DATA,
        ["Index", "Table", "Unique", "Columns", "Partial condition", "Purpose"],
        index_rows,
        [26, 28, 12, 48, 42, 55],
    )

    integrity_rows = [
        ["Ledger cannot be edited", "xp_ledger", "Trigger trg_xl_no_update, trg_xl_no_delete. Reversal rows must point at an award (CHECK)."],
        ["Journal cannot be edited", "goal_contribution", "Trigger trg_gc_no_update, trg_gc_no_delete."],
        ["Audit cannot be edited", "audit_log", "Trigger trg_audit_no_update, trg_audit_no_delete, plus the hash chain (prev_hash, hash)."],
        ["Financial fields on the master are frozen", "txn_master", "Trigger trg_txn_master_immutable. Amount may change only while PENDING. Trigger trg_txn_master_no_delete."],
        ["One live version per rule", "rule_version", "Partial unique index ux_rv_one_live WHERE status IN ('ACTIVE','PAUSED')."],
        ["No duplicate delivery", "ingestion_event", "UNIQUE (source_system, source_event_id)."],
        ["No duplicate transaction", "txn_master", "UNIQUE (source_system, source_txn_id)."],
        ["No duplicate reward", "xp_ledger", "UNIQUE idempotency_key."],
        ["No duplicate nudge", "nudge", "UNIQUE dedupe_key."],
        ["No duplicate daily balance", "balance_snapshot", "UNIQUE (account_id, as_of_date)."],
        ["Referential integrity", "All", "Foreign keys are on. Reference-data foreign keys are enforced."],
        ["Refund cannot exceed amount", "txn_master", "CHECK refunded_minor <= amount_minor and amount_minor >= 0."],
        ["Ledger sign", "xp_ledger", "AWARD points > 0 and reverses_ledger_id is null. REVERSAL points < 0 and reverses_ledger_id is set. ADJUSTMENT is either sign."],
        ["Campaign window", "campaign", "CHECK end_date >= start_date."],
        ["One open version", "rule_version", "Application rule OPEN_VERSION_EXISTS: only one draft, pending or approved version at a time, plus one live version."],
    ]
    for table in table_names:
        for check in extract_checks(master.get(table, "")):
            integrity_rows.append(["CHECK", table, check])
    for name, tbl, sql in conn.execute(
        "SELECT name, tbl_name, sql FROM sqlite_master WHERE type='trigger' ORDER BY tbl_name, name"
    ):
        action = "UPDATE and DELETE blocked" if "no_update" in name or "no_delete" in name or "immutable" in name else "See definition"
        msg = ""
        mm = re.search(r"RAISE\(ABORT,\s*'([^']+)'\)", sql)
        if mm:
            msg = mm.group(1)
        integrity_rows.append([f"Trigger {name}", tbl, msg or action])
    add_sheet(
        wb,
        "Integrity Rules",
        DATA,
        ["Guarantee", "Table", "Mechanism"],
        integrity_rows,
        [48, 28, 100],
    )

    enum_rows = []
    for table in table_names:
        for check in extract_checks(master.get(table, "")):
            m = re.match(r"(\w+)\s+IN\s*\((.*)\)\s*$", check, re.S)
            if not m:
                continue
            col, raw = m.group(1), m.group(2)
            values = re.findall(r"'([^']*)'|(\d+)", raw)
            shown = ", ".join(a or b for a, b in values)
            enum_rows.append([table, col, shown])
    add_sheet(
        wb,
        "Allowed Values",
        DATA,
        ["Table", "Column", "Allowed values"],
        enum_rows,
        [32, 28, 110],
    )

    add_sheet(
        wb,
        "Postgres Mapping",
        DATA,
        ["SQLite", "PostgreSQL"],
        [
            ["TEXT ISO timestamps", "timestamptz (store UTC). Business dates as date."],
            ["INTEGER money", "bigint (minor units)"],
            ["INTEGER PRIMARY KEY AUTOINCREMENT", "bigint GENERATED ALWAYS AS IDENTITY"],
            ["INTEGER 0/1 flags", "boolean"],
            ["*_json TEXT", "jsonb (add GIN indexes if you query attributes)"],
            ["RAISE(ABORT, …) triggers", "BEFORE UPDATE OR DELETE trigger functions that RAISE EXCEPTION. Also REVOKE UPDATE, DELETE from the application role on the append-only tables."],
            ["CHECK / partial indexes", "Supported as written"],
            ["BEGIN / SAVEPOINT in Db.tx", "Same semantics"],
            ["Table partitioning", "Partition txn_master, rule_evaluation, goal_contribution and audit_log by month when volumes warrant"],
            ["db.raw single-writer", "Use a pool. Keep one transaction per event. The worker takes events with FOR UPDATE SKIP LOCKED so several workers can run."],
        ],
        [42, 110],
    )

    add_sheet(
        wb,
        "Retention",
        DATA,
        ["Data", "Hot", "Archive"],
        [
            ["txn_master, xp_ledger, audit_log", "Policy-driven (typically 7+ years)", "Partition then cold storage. Never delete ledger rows."],
            ["ingestion_event.payload_json", "90 days", "After that, null the payload and keep the hash."],
            ["rule_evaluation.trace_json", "90 days", "Set storeTrace to false for very high volumes. Aggregate outcome rows stay."],
            ["nudge", "12 months", "Delete"],
        ],
        [42, 42, 78],
    )

    add_sheet(
        wb,
        "Rule Definition",
        RULE,
        ["Field", "Required", "Type / allowed values", "Default", "Meaning"],
        [
            ["name", "Yes", "string, 3–120", "", "Display name"],
            ["description", "No", "string, max 1000", "empty", "Plain description"],
            ["businessObjective", "No", "string, max 1000", "empty", "Why the goal exists"],
            ["category", "Yes", "SPEND, FREQUENCY, BILL_PAYMENT, BALANCE, SAVINGS_GROWTH, COMBINED, ENGAGEMENT, OTHER", "", "Business category"],
            ["campaignCode", "No", "string or null", "null", "Campaign the rule belongs to"],
            ["kind", "Yes", "GOAL, PER_EVENT, COMBINED", "", "What the rule computes"],
            ["scope.channels", "No", "list of channel codes", "empty = all", "Cheap pre-filter before conditions run"],
            ["scope.productCodes", "No", "list of product codes", "empty = all", "Cheap pre-filter before conditions run"],
            ["eligibility", "No", "condition group", "empty AND group", "Who can earn"],
            ["qualifying", "No", "condition group", "empty AND group", "What counts toward the metric"],
            ["metric.type", "GOAL only", "COUNT, SUM, AVG_BALANCE, MIN_BALANCE, BALANCE_GROWTH", "", "How progress is measured"],
            ["metric.operator", "No", "GTE, GT", "GTE", "How a tier threshold is compared"],
            ["period.type", "Yes", "DAILY, WEEKLY, MONTHLY, QUARTERLY, CAMPAIGN", "", "Window the metric is computed over"],
            ["period.timezone", "No", "IANA time zone", "Asia/Kolkata", "Period boundaries are computed in this zone"],
            ["period.dateBasis", "No", "TXN_DATE, POSTING_DATE", "TXN_DATE", "Which date places a transaction in a period"],
            ["period.lateGraceDays", "No", "integer 0–31", "0", "Days after period end before it locks"],
            ["tiers[].label", "No", "string, max 60", "", "Tier name"],
            ["tiers[].threshold", "Yes if tier", "number", "", "Rupees for money metrics, a count for COUNT and COMBINED"],
            ["tiers[].points", "Yes if tier", "integer ≥ 0", "", "XP for that tier"],
            ["tierMode", "No", "CUMULATIVE, HIGHEST_ONLY", "CUMULATIVE", "Each tier adds, or only the highest tier counts"],
            ["perEvent.mode", "PER_EVENT", "FIXED, PER_AMOUNT", "", "Fixed points, or points per amount"],
            ["perEvent.points", "PER_EVENT", "positive integer", "", "Points, or points per chunk"],
            ["perEvent.perAmount", "PER_AMOUNT", "positive number", "", "Rupees per points chunk. 100 means 1 point per ₹100 when points is 1"],
            ["combined.components[].ruleCode", "COMBINED", "existing rule code", "", "Component goal"],
            ["combined.components[].minTier", "No", "integer ≥ 1", "1", "Component counts as met at this tier or higher"],
            ["balancePolicy.gapFill", "No", "NONE, CARRY_FORWARD", "NONE", "NONE refuses to guess a missing snapshot"],
            ["limits.frequency", "No", "ONCE, RECURRING", "RECURRING", "Whether the goal can pay more than once"],
            ["limits.maxOccurrences", "No", "positive integer or null", "null", "Cap on how many times it can pay"],
            ["limits.capPointsPerPeriod", "No", "integer ≥ 0 or null", "null", "Points cap inside one period"],
            ["limits.capPointsLifetime", "No", "integer ≥ 0 or null", "null", "Points cap across all periods"],
            ["handling.reversal", "No", "CLAWBACK, GRANDFATHER", "CLAWBACK", "Whether a drop caused by a reversal withdraws points"],
            ["handling.countStatuses", "No", "POSTED, and optionally PENDING", "POSTED", "Which statuses count. FAILED never counts"],
            ["priority", "No", "integer 0–1000", "100", "Higher number wins an EXCLUSIVE_PRIORITY group"],
            ["conflict.group", "No", "string max 40, or null", "null", "Rules in the same group are resolved together"],
            ["conflict.policy", "No", "CUMULATIVE, EXCLUSIVE_HIGHEST_POINTS, EXCLUSIVE_PRIORITY", "CUMULATIVE", "How members of the group share a period"],
            ["effective.from", "No", "YYYY-MM-DD or null", "null", "First date that counts"],
            ["effective.to", "No", "YYYY-MM-DD or null", "null", "Last date that counts"],
            ["nudges.enabled", "No", "boolean", "true", "Whether this rule generates nudges"],
            ["nudges.nearTargetPct", "No", "number 1–99", "80", "Progress percent that raises NEAR_TARGET"],
            ["backfill", "No", "boolean", "false", "On activation, replay the open period from the master"],
        ],
        [40, 16, 62, 22, 62],
    )

    add_sheet(
        wb,
        "Condition Fields",
        RULE,
        ["Scope", "Field", "Label", "Type", "Allowed values / source"],
        [
            ["txn", "txn.type", "Transaction type", "enum", "txn_type_ref"],
            ["txn", "txn.category", "Transaction category", "enum", "txn_category_ref"],
            ["txn", "txn.channel", "Channel", "enum", "channel_ref"],
            ["txn", "txn.product_code", "Product code", "enum", "product_ref"],
            ["txn", "txn.instrument_type", "Instrument type", "enum", "CREDIT_CARD, DEBIT_CARD, UPI_VPA, WALLET"],
            ["txn", "txn.amount", "Transaction amount (INR)", "money", "Compared in minor units. txn.amount GTE 10000 is ₹10,000.00 exactly, so ₹9,999.99 never qualifies."],
            ["txn", "txn.status", "Transaction status", "enum", "PENDING, POSTED, REVERSED, REFUNDED, PARTIALLY_REFUNDED, FAILED"],
            ["txn", "txn.dr_cr", "Debit / credit", "enum", "D, C"],
            ["txn", "txn.merchant_id", "Merchant ID", "string", ""],
            ["txn", "txn.merchant_name", "Merchant name", "string", ""],
            ["txn", "txn.mcc", "Merchant category code", "string", ""],
            ["txn", "txn.day_of_week", "Day of week", "number", "1 = Monday … 7 = Sunday, in the rule time zone"],
            ["txn", "txn.day_of_month", "Day of month", "number", ""],
            ["txn", "txn.hour_of_day", "Hour of day", "number", "0–23 in the rule time zone"],
            ["account", "account.product_code", "Account product code", "enum", "product_ref"],
            ["account", "account.product_family", "Account product family", "enum", "SAVINGS, CURRENT, CREDIT_CARD, DEBIT_CARD, UPI, LOAN, DEPOSIT, WALLET, OTHER"],
            ["account", "account.type", "Account type", "enum", "SAVINGS, CURRENT, CREDIT_CARD, DEBIT_CARD, UPI, LOAN, DEPOSIT, WALLET, OTHER"],
            ["account", "account.status", "Account status", "enum", "ACTIVE, DORMANT, FROZEN, CLOSED"],
            ["account", "account.age_days", "Account age (days)", "number", ""],
            ["customer", "customer.segments", "Customer segments", "array", "segment reference data, honouring valid_from / valid_to"],
            ["customer", "customer.type", "Customer type", "enum", "RETAIL, SME, CORPORATE"],
            ["customer", "customer.kyc_status", "KYC status", "enum", "FULL, MINIMAL, PENDING"],
            ["customer", "customer.status", "Customer status", "enum", "ACTIVE, INACTIVE, CLOSED, BLOCKED"],
            ["txn.attr", "txn.attr.<key>", "Extensible attribute", "string", "Any attribute the source sent. No schema change."],
        ],
        [14, 28, 32, 12, 88],
    )

    add_sheet(
        wb,
        "Operators",
        RULE,
        ["Operator", "Meaning", "Missing value"],
        [
            ["EQ", "Equals", "Fails when the value is missing, so an unknown value never qualifies"],
            ["NE", "Does not equal", "Passes when the value is missing"],
            ["GT", "Greater than", "Fails when the value is missing"],
            ["GTE", "Greater than or equal", "Fails when the value is missing"],
            ["LT", "Less than", "Fails when the value is missing"],
            ["LTE", "Less than or equal", "Fails when the value is missing"],
            ["IN", "Is one of the listed values", "Fails when the value is missing"],
            ["NOT_IN", "Is not one of the listed values", "Passes when the value is missing"],
            ["BETWEEN", "Is between the two ends, inclusive", "Fails when the value is missing"],
            ["CONTAINS", "Contains the text", "Fails when the value is missing"],
            ["STARTS_WITH", "Starts with the text", "Fails when the value is missing"],
            ["IS_NULL", "Is empty", "Passes when the value is missing"],
            ["NOT_NULL", "Is present", "Fails when the value is missing"],
        ],
        [16, 48, 70],
    )

    add_sheet(
        wb,
        "Evaluation Logic",
        RULE,
        ["Step", "Rule"],
        [
            ["1. Scope", "A transaction is offered only to live rule versions. A cheap channel and product pre-filter runs first. Out of scope produces no evaluation row."],
            ["2. Eligibility", "The eligibility group must be true: customer active, KYC, segments, effective dates, campaign window and enrolment. Otherwise the outcome is INELIGIBLE and the trace explains which condition failed."],
            ["3. Qualifying", "The status must be in countStatuses and the qualifying group must be true. Otherwise the outcome is NOT_QUALIFYING."],
            ["4. Journal", "A qualifying transaction appends APPLY, or REVERSE, to goal_contribution with amount and count deltas for the period. The journal is the source of truth for progress."],
            ["5. Progress", "goal_progress for (customer, rule version, period) is updated: txn_count, txn_sum_minor, metric_value, achieved_tier. No query reads other transactions."],
            ["6. Slots", "Each rule, customer and period has slots: tier:1, tier:2, …, best (HIGHEST_ONLY), or txn:<id> for a per-event rule. The engine computes the points that slot should currently be worth."],
            ["7. Compare", "target > net posts an AWARD for the difference. target < net posts a REVERSAL only when the rule is CLAWBACK and the drop is caused by a reversal or refund. GRANDFATHER, or a drop that is not reversal-driven, keeps the points and records NO_CHANGE. target = net posts nothing."],
            ["8. Closed period", "Late data for a closed period updates the metric but cannot earn. Outcome PERIOD_CLOSED."],
            ["9. Frequency", "ONCE and maxOccurrences block further awards after the limit."],
            ["10. Conflicts", "EXCLUSIVE_HIGHEST_POINTS: most points wins, then priority. EXCLUSIVE_PRIORITY: higher priority number wins, then rule code. A group member that already rewarded the customer for the period blocks the others. CUMULATIVE stacks."],
            ["11. Caps and budget", "Applied in deterministic order, higher priority first: per period, lifetime, then campaign points_budget. Points are trimmed, not dropped, when part of a cap remains. Outcome CAPPED when a cap binds."],
            ["12. Idempotency", "The result depends only on the facts. Re-processing any event, in any order, any number of times, converges to the same ledger. The idempotency key (AW|… / RV|…) is a second line of defence."],
            ["13. Balance goals", "Evaluated when a snapshot arrives and at period close, using the same resolution step. Balance metrics never come from transactions."],
            ["14. Missing snapshots", "If a required day is missing and gapFill is NONE, progress is DATA_UNAVAILABLE and no reward is issued. An in-period average is provisional until the last snapshot arrives or the period ends."],
            ["15. Combined goals", "Re-evaluated whenever a component's achieved tier changes. The metric is the number of components whose achieved tier is at least minTier."],
            ["16. Time", "DAILY, WEEKLY, MONTHLY and QUARTERLY periods are computed in the rule time zone. CAMPAIGN uses the campaign window. 23:59 IST on the last day and 00:00 the next day are different months, regardless of UTC."],
            ["17. Period close", "A period closes when today (rule time zone) is past period_end + lateGraceDays. Progress becomes CLOSED_ACHIEVED or CLOSED_NOT_ACHIEVED. Later arrivals update the metric only."],
            ["18. Effective dates", "A transaction outside effective.from .. effective.to, or outside the campaign window, never counts."],
            ["19. Reversals and refunds", "They arrive linked by originalSourceTxnId, or as a status change on the original. The journal gets a REVERSE entry, progress falls, and points are clawed back only if the new target is lower and the rule says CLAWBACK."],
            ["20. Status", "PENDING counts only if the rule lists it. FAILED never counts. Status moves forward only. An older status arriving late is ignored and recorded."],
            ["21. Versions", "Only one version per rule may be open (draft, pending, approved) and one live (ACTIVE or PAUSED). Editing is allowed only in DRAFT, by its owner. A SYSTEM_ADMIN may edit another maker's draft; the edit is audited."],
            ["22. Activation", "Activation can backfill the open period from the master. Without backfill, a new version counts only events arriving after activation. Pause and resume do not catch up; run backfill afterwards."],
            ["23. History", "Goal progress and ledger rows keep the rule_version_id that produced them, so history stays explainable after later changes."],
            ["24. Explainability", "Every evaluation stores a trace: each condition with the value seen, the operator, the expected value, and pass or fail; the metric before and after; the tier targets; and the reason for any block."],
            ["25. Validation", "Unknown fields, bad operators, inverted ranges, unsorted tiers and unknown component rules are errors. Values not found in reference data are warnings."],
            ["26. Sealed definition", "The JSON definition is hashed at submit and verified on every load. A mismatch raises DEFINITION_TAMPERED."],
            ["27. Maker-checker", "The approver must be a different person from the creator and the submitter. Approval is by identity, not only by role."],
            ["28. Money boundary", "Conditions compare money in minor units. ₹9,999.99 does not satisfy a ₹10,000 threshold."],
        ],
        [22, 130],
    )

    add_sheet(
        wb,
        "Metrics",
        RULE,
        ["Metric", "Used by", "Computation", "Reward if data is missing"],
        [
            ["COUNT", "GOAL", "Number of qualifying transactions in the period. A refund does not reduce a counted transaction unless it is fully reversed.", "Not a balance metric"],
            ["SUM", "GOAL", "Qualifying amount minus refunded amount.", "Not a balance metric"],
            ["AVG_BALANCE", "GOAL", "Mean, over each day in the period, of the sum of closing balances of qualifying accounts, from balance_snapshot.", "DATA_UNAVAILABLE and no reward when gapFill is NONE"],
            ["MIN_BALANCE", "GOAL", "Lowest such daily total in the period.", "DATA_UNAVAILABLE and no reward when gapFill is NONE"],
            ["BALANCE_GROWTH", "GOAL", "Last day's total minus the total on the day before the period.", "DATA_UNAVAILABLE and no reward when gapFill is NONE"],
            ["Components met", "COMBINED", "Number of component goals whose achieved tier is at least minTier. A tier threshold is that count (both goals = 2).", "Depends on each component"],
            ["FIXED", "PER_EVENT", "A fixed number of points for each qualifying transaction. Slot is txn:<id>.", "Not a balance metric"],
            ["PER_AMOUNT", "PER_EVENT", "points for every perAmount rupees. Example: 1 XP per ₹100.", "Not a balance metric"],
        ],
        [20, 16, 100, 55],
    )

    add_sheet(
        wb,
        "Award Resolution",
        RULE,
        ["Order", "Check", "What happens"],
        [
            ["1", "Net and delta per slot", "Compare the slot target with the net already in the ledger."],
            ["2", "Downward moves", "CLAWBACK posts a reversal when a reversal or refund lowered the target. GRANDFATHER keeps the points and records NO_CHANGE."],
            ["3", "Closed-period lock", "Late data for a closed period updates the metric and cannot earn."],
            ["4", "Frequency", "ONCE and maxOccurrences stop further awards."],
            ["5", "Conflict group", "EXCLUSIVE_HIGHEST_POINTS or EXCLUSIVE_PRIORITY picks one winner. CUMULATIVE allows all. A member that already paid this period blocks the others."],
            ["6", "Caps and budget", "Per period, then lifetime, then campaign points_budget. Higher priority is applied first. A partial cap trims the points."],
            ["7", "Post", "Write AWARD or REVERSAL with an idempotency key, then the evaluation row with an outcome and a reason."],
        ],
        [10, 28, 110],
    )

    add_sheet(
        wb,
        "Lifecycle",
        RULE,
        ["From", "To", "Action", "Rule"],
        [
            ["—", "DRAFT", "Create", "Editing is allowed only in DRAFT, by its owner."],
            ["DRAFT", "PENDING_APPROVAL", "Submit", "Validation must pass. The definition is sealed with a hash."],
            ["PENDING_APPROVAL", "APPROVED", "Approve", "Approver must not be the creator or the submitter."],
            ["PENDING_APPROVAL", "REJECTED", "Reject", "A reason is required."],
            ["REJECTED", "DRAFT", "New version", "A new version is created from the rejected one."],
            ["APPROVED", "ACTIVE", "Activate", "Optional backfill of the open period. The previous live version becomes SUPERSEDED."],
            ["ACTIVE", "PAUSED", "Pause", "A reason is required. Pause does not catch up later by itself."],
            ["PAUSED", "ACTIVE", "Resume", "Resume does not replay missed events. Run backfill afterwards if catch-up is needed."],
            ["ACTIVE", "SUPERSEDED", "Newer version activates", "History keeps the version that produced each ledger row."],
            ["ACTIVE", "RETIRED", "Retire", "The rule stops evaluating new events."],
            ["PAUSED", "RETIRED", "Retire", "The rule stops evaluating new events."],
            ["SUPERSEDED", "ACTIVE", "Rollback", "A reason of at least 5 characters is required."],
        ],
        [22, 22, 28, 78],
    )

    campaigns = json.loads(CAMPAIGNS.read_text())["campaigns"]
    add_sheet(
        wb,
        "Campaigns",
        RULE,
        ["Code", "Name", "Description", "Objective", "Start", "End", "Points budget", "Sample load"],
        [
            [
                c["code"],
                c["name"],
                c["description"],
                c["objective"],
                c["startDate"],
                c["endDate"],
                c.get("pointsBudget", ""),
                "Created and activated by the sample loader",
            ]
            for c in campaigns
        ],
        [22, 28, 52, 32, 14, 14, 16, 42],
    )

    sample_rows = []
    condition_rows = []
    tier_rows = []
    rule_order = [
        "SAV_AVG_BAL_50K",
        "CARD_SPEND_10K",
        "UPI_10_TXNS",
        "UTILITY_3_BILLS",
        "SAV_GROWTH_25K",
        "BALANCE_PLUS_CARD",
        "UPI_MERCHANT_PER_TXN",
    ]
    extra = sorted(p.stem for p in RULES_DIR.glob("*.json") if p.stem not in rule_order)
    for code in rule_order + extra:
        path = RULES_DIR / f"{code}.json"
        doc = json.loads(path.read_text())
        d = doc["definition"]
        period = d.get("period") or {}
        limits = d.get("limits") or {}
        handling = d.get("handling") or {}
        conflict = d.get("conflict") or {}
        effective = d.get("effective") or {}
        nudges = d.get("nudges") or {}
        balance = d.get("balancePolicy") or {}
        scope = d.get("scope") or {}
        metric = d.get("metric") or {}
        per = d.get("perEvent") or {}
        combined = (d.get("combined") or {}).get("components") or []
        omitted = []
        for key, default in [
            ("tierMode", "CUMULATIVE"),
            ("limits.frequency", "RECURRING"),
            ("limits.maxOccurrences", "none"),
            ("limits.capPointsPerPeriod", "none"),
            ("limits.capPointsLifetime", "none"),
            ("handling.reversal", "CLAWBACK"),
            ("handling.countStatuses", "POSTED"),
            ("conflict.group", "none"),
            ("conflict.policy", "CUMULATIVE"),
            ("period.dateBasis", "TXN_DATE"),
            ("period.lateGraceDays", "0"),
            ("balancePolicy.gapFill", "NONE"),
            ("nudges.enabled", "true"),
            ("nudges.nearTargetPct", "80"),
            ("backfill", "false"),
            ("priority", "100"),
        ]:
            parent, _, child = key.partition(".")
            present = child in d.get(parent, {}) if child else key in d
            if child and parent not in d:
                present = False
            if not child and key not in d:
                present = False
            if not present:
                omitted.append(f"{key}={default}")
        tiers = d.get("tiers") or []
        tier_text = "; ".join(
            f"{t.get('label', '')}: {t['threshold']} → {t['points']} XP" for t in tiers
        )
        component_text = "; ".join(f"{c['ruleCode']} tier ≥ {c.get('minTier', 1)}" for c in combined)
        load_status = (
            "Created and submitted; left PENDING_APPROVAL so the queue is not empty"
            if doc["code"] == "UPI_MERCHANT_PER_TXN"
            else "Created, submitted, approved and activated"
        )
        sample_rows.append(
            [
                doc["code"],
                d.get("name", ""),
                d.get("description", ""),
                d.get("businessObjective", ""),
                d.get("category", ""),
                d.get("campaignCode", ""),
                d.get("kind", ""),
                ", ".join(scope.get("channels") or []) or "all",
                ", ".join(scope.get("productCodes") or []) or "all",
                metric.get("type", ""),
                metric.get("operator", "GTE") if metric else "",
                period.get("type", ""),
                period.get("timezone", "Asia/Kolkata"),
                period.get("dateBasis", "TXN_DATE"),
                period.get("lateGraceDays", 0),
                d.get("tierMode", "CUMULATIVE"),
                tier_text,
                per.get("mode", ""),
                per.get("points", ""),
                per.get("perAmount", ""),
                component_text,
                limits.get("frequency", "RECURRING"),
                "" if limits.get("maxOccurrences") is None else limits.get("maxOccurrences", ""),
                "" if "capPointsPerPeriod" not in limits else ("" if limits.get("capPointsPerPeriod") is None else limits.get("capPointsPerPeriod")),
                handling.get("reversal", "CLAWBACK"),
                ", ".join(handling.get("countStatuses") or ["POSTED"]),
                d.get("priority", 100),
                conflict.get("group") or "",
                conflict.get("policy", "CUMULATIVE"),
                effective.get("from") or "",
                effective.get("to") or "",
                balance.get("gapFill", "NONE"),
                "Yes" if nudges.get("enabled", True) else "No",
                nudges.get("nearTargetPct", 80),
                "Yes" if d.get("backfill", False) else "No",
                doc.get("changeNote", ""),
                load_status,
                "; ".join(omitted),
            ]
        )
        for seq, cond in enumerate(flatten_conditions(d), 1):
            condition_rows.append([doc["code"], seq, *cond])
        for i, t in enumerate(tiers, 1):
            tier_rows.append([doc["code"], "Tier", i, t.get("label", ""), t["threshold"], t["points"], "", ""])
        for c in combined:
            tier_rows.append([doc["code"], "Component", "", "", "", "", c["ruleCode"], c.get("minTier", 1)])

    add_sheet(
        wb,
        "Sample Rules",
        RULE,
        [
            "Code",
            "Name",
            "Description",
            "Business objective",
            "Category",
            "Campaign",
            "Kind",
            "Scope channels",
            "Scope products",
            "Metric",
            "Metric operator",
            "Period",
            "Time zone",
            "Date basis",
            "Late grace days",
            "Tier mode",
            "Tiers",
            "Per-event mode",
            "Per-event points",
            "Per amount (INR)",
            "Components",
            "Frequency",
            "Max occurrences",
            "Cap points per period",
            "Reversal",
            "Count statuses",
            "Priority",
            "Conflict group",
            "Conflict policy",
            "Effective from",
            "Effective to",
            "Balance gap fill",
            "Nudges",
            "Near-target %",
            "Backfill",
            "Change note",
            "Sample load status",
            "Engine defaults (not set in the JSON file)",
        ],
        sample_rows,
        [24, 42, 55, 32, 18, 18, 14, 22, 16, 18, 16, 14, 16, 16, 16, 16, 42, 16, 16, 16, 42, 14, 16, 20, 14, 16, 12, 16, 26, 16, 14, 16, 12, 14, 12, 18, 55, 55],
    )
    add_sheet(
        wb,
        "Rule Conditions",
        RULE,
        ["Rule", "Seq", "Group", "Nesting", "Negated", "Field", "Operator", "Value", "Plain language"],
        condition_rows,
        [24, 8, 14, 28, 12, 28, 14, 22, 70],
    )
    add_sheet(
        wb,
        "Tiers and Components",
        RULE,
        ["Rule", "Row type", "Tier no", "Label", "Threshold", "Points", "Component rule", "Minimum tier"],
        tier_rows,
        [24, 14, 12, 20, 14, 12, 22, 14],
    )

    add_sheet(
        wb,
        "Sample Outcomes",
        RULE,
        ["Customer", "Profile / facts", "Rule", "Result"],
        [
            ["CUST-1001", "Affluent. Card ₹4,200 + ₹5,100 + ₹3,100 = ₹12,400", "CARD_SPEND_10K", "Tier 1 (₹10,000) → 300 XP"],
            ["CUST-1001", "12 UPI merchant payments", "UPI_10_TXNS", "100 XP"],
            ["CUST-1001", "3 utility bills", "UTILITY_3_BILLS", "150 XP"],
            ["CUST-1001", "Savings about ₹62,000–₹66,000 all of last month", "SAV_AVG_BAL_50K", "500 XP (September)"],
            ["CUST-1001", "One transaction delivered 4 times", "Any live rule", "Exactly one award. Duplicates do not pay again."],
            ["CUST-1002", "Salaried. Card ₹6,999 + ₹3,000 = ₹9,999", "CARD_SPEND_10K", "One rupee short → 0 XP. NEAR_TARGET nudge."],
            ["CUST-1002", "9 UPI payments + 1 self-transfer", "UPI_10_TXNS", "Self-transfer excluded → 9 < 10 → 0 XP"],
            ["CUST-1002", "Average savings about ₹48,500", "SAV_AVG_BAL_50K", "Just misses → 0 XP. Progress and near-target nudges."],
            ["CUST-1003", "Digital-first. Savings rise from ₹31,000 to ₹61,000", "SAV_GROWTH_25K", "Growth ₹30,000 → 400 XP"],
            ["CUST-1003", "3 utility bills across the campaign (1 last month, 2 this month)", "UTILITY_3_BILLS", "150 XP"],
            ["CUST-1004", "Affluent. Savings flat at ₹55,000", "SAV_AVG_BAL_50K", "500 XP"],
            ["CUST-1004", "Card spend ₹10,500 last month", "CARD_SPEND_10K", "300 XP (September)"],
            ["CUST-1004", "Balance ₹55,000 and card ₹10,500 last month", "BALANCE_PLUS_CARD", "Both components met → 1,000 XP"],
            ["CUST-1004", "This month ₹8,000 + ₹3,000 + ₹1,500, then a ₹3,000 refund (₹12,500 then ₹9,500)", "CARD_SPEND_10K", "300 awarded, then ₹9,500 < ₹10,000 → −300 clawback"],
            ["CUST-1005", "New customer, no nudge consent. 6 small UPI payments. Balance ₹8,000", "All goals", "0 XP. Nudges suppressed with reason NO_CONSENT."],
            ["UPI_MERCHANT_PER_TXN", "Sample rule is left pending approval", "UPI_MERCHANT_PER_TXN", "Does not award in the sample load, because it is not active."],
            ["All customers", "After the sample load", "Ledger totals", "Awarded 3,700 XP. Reversed −300 XP. Net 3,400 XP."],
            ["All customers", "244 ingestion events", "Ingestion", "244 processed, 0 rejected, 0 duplicate awards. Reconciliation PASSED. Audit chain valid."],
        ],
        [22, 78, 28, 70],
    )

    # Contents
    cover.sheet_properties.tabColor = NAVY
    cover.sheet_view.showGridLines = False
    cover.page_setup.orientation = "landscape"
    cover.page_setup.fitToPage = True
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.page_setup.paperSize = cover.PAPERSIZE_A4
    cover.oddFooter.right.text = "Page &P of &N"
    cover.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.5)
    cover.merge_cells("A1:D1")
    cover["A1"] = "Money XP (MXP)"
    cover["A1"].font = TITLE_FONT
    cover.merge_cells("A2:D2")
    cover["A2"] = "Requirements, table design and rules"
    cover["A2"].font = Font(name="Calibri", size=14, color="3D4F5F")
    cover.row_dimensions[1].height = 28
    cover.row_dimensions[2].height = 20

    intro = (
        "Money XP rewards customers with XP points for banking behaviour already visible in transaction data: "
        "card spend, UPI payments, bill payments, and keeping or growing balances. Business teams configure goals as data. "
        "The platform evaluates them against a central transaction master, posts rewards to an immutable ledger, and nudges "
        "customers toward the next goal. MXP never posts to or alters financial transactions. It reads copies that source "
        "systems send, and writes only to its own tables."
    )
    cover.merge_cells("A4:D4")
    cover["A4"] = intro
    cover["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    cover["A4"].font = Font(name="Calibri", size=11)
    cover.row_dimensions[4].height = 62

    meta = [
        ("Phase", "Phase 1"),
        ("Priority key", "M = must for Phase 1. S = should. Only F-42 is a should."),
        ("Sources", "docs/01-requirements.md, docs/03-data-model.md, docs/04-rules-and-evaluation.md, docs/08-sample-data.md, packages/engine/src/schema.sql, samples/rules, samples/campaigns.json"),
        ("Acceptance", "An admin creates a goal, configures conditions and XP, tests it against scenarios, has it approved and activated, and verifies correct rewards with no duplicates."),
    ]
    row = 6
    for label, value in meta:
        cover.cell(row, 1, label).font = LABEL_FONT
        cover.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        cover.cell(row, 2, value).font = BODY
        cover.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        cover.row_dimensions[row].height = 32 if label in ("Sources", "Acceptance") else 18
        row += 1

    cover.cell(11, 1, "Sheets").font = SECTION_FONT
    headers = ["Sheet", "What it contains"]
    for c, h in enumerate(headers, 1):
        cell = cover.cell(12, c, h)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    catalog = [
        ("Actors", "Who uses MXP and the six console roles"),
        ("Use Cases", "The seven configurable business goals and their sample rule codes"),
        ("Functional Requirements", "F-01 through F-72, with area and must / should priority"),
        ("Non-Functional", "Correctness, performance, security, audit, availability, privacy"),
        ("Out of Scope", "What Phase 1 does not include, and the seam left for later"),
        ("Conventions", "Money, time, facts versus projections, privacy, acceptance"),
        ("Table Catalog", "Every table, its domain, and whether it is a fact, projection or master"),
        ("Column Design", "Every column: type, key, required, default, foreign key, check, comment"),
        ("Relationships", "Foreign keys and what they mean"),
        ("Indexes", "Every index, including the one-live-version partial unique index"),
        ("Integrity Rules", "Triggers, checks and uniqueness guarantees"),
        ("Allowed Values", "Every CHECK … IN list (statuses, kinds, roles, outcomes)"),
        ("Postgres Mapping", "How the SQLite schema maps when you move to PostgreSQL"),
        ("Retention", "How long each store is kept"),
        ("Rule Definition", "Every field of a rule, with type and engine default"),
        ("Condition Fields", "txn, account, customer and extensible attribute fields"),
        ("Operators", "The 13 operators and how a missing value is treated"),
        ("Evaluation Logic", "Scope, eligibility, journal, slots, caps, conflicts, reversals, versions"),
        ("Metrics", "COUNT, SUM, balance metrics, combined goals and per-event modes"),
        ("Award Resolution", "The order of checks that posts an award or a reversal"),
        ("Lifecycle", "Draft through approve, activate, pause, retire and rollback"),
        ("Campaigns", "MXP_ALWAYS_ON and UTILITY_FEST"),
        ("Sample Rules", "All seven shipped rules with effective values and engine defaults"),
        ("Rule Conditions", "Every eligibility and qualifying condition in plain language"),
        ("Tiers and Components", "Points ladders and combined-goal components"),
        ("Sample Outcomes", "What each sample customer earns, including the clawback and the ledger totals"),
    ]
    for i, (name, desc) in enumerate(catalog):
        r = 13 + i
        cell = cover.cell(r, 1, name)
        cell.font = LINK_FONT
        cell.hyperlink = f"#'{name}'!A1"
        cover.cell(r, 2, desc).font = BODY
        if i % 2 == 1:
            for c in (1, 2):
                if c == 1:
                    continue
                cover.cell(r, c).fill = PatternFill("solid", fgColor=ALT)
        cover.row_dimensions[r].height = 18
    cover.column_dimensions["A"].width = 32
    cover.column_dimensions["B"].width = 88
    cover.column_dimensions["C"].width = 18
    cover.column_dimensions["D"].width = 18
    cover.freeze_panes = "A13"
    cover.oddHeader.left.text = "Money XP (MXP)"
    cover.print_title_rows = "12:12"
    cover.page_setup.fitToHeight = 0
    cover.sheet_properties.pageSetUpPr.fitToPage = True

    # Counts for a small summary under the intro area — placed at column C of meta? Keep contents clean.
    wb.properties.title = "MXP requirements, table design and rules"
    wb.properties.subject = "Phase 1 functional requirements, schema and evaluation rules"
    wb.properties.creator = "MXP"

    # Move contents is already first.
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Sheets: {len(wb.sheetnames)}")
    print(f"Columns: {len(column_rows)}  Requirements: {len(functional)}  Rules: {len(sample_rows)}  Conditions: {len(condition_rows)}")


if __name__ == "__main__":
    main()
